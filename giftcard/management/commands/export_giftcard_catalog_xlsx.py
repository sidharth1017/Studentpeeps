import csv
import json
from decimal import Decimal
from pathlib import Path

from django.core.management.base import BaseCommand
from django.utils import timezone

from giftcard.models import ProductOverride, ProviderProduct

DEFAULT_RATES_CSV = Path(
    "/Users/sidharthverma/.cursor/projects/"
    "Users-sidharthverma-Studentpeeps-codebase-Studentpeeps/attachments/"
    "61931cf4-5595-4e4d-9fc5-2dd7b5b77532/Discount_Sheet_-_Rates.csv"
)


def _cell(value):
    if value is None:
        return ""
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, bool):
        return value
    if isinstance(value, (list, dict)):
        return json.dumps(value, ensure_ascii=False)
    return str(value)


def _product_row(p: ProviderProduct) -> dict:
    return {
        "id": p.pk,
        "provider_id": p.provider_id,
        "provider_name": p.provider.name if p.provider_id else "",
        "sku": p.sku,
        "provider_product_id": p.provider_product_id,
        "name": p.name,
        "url": p.url,
        "description": p.description,
        "short_description": p.short_description,
        "brand_name": p.brand_name,
        "brand_code": p.brand_code,
        "brand_logo": p.brand_logo,
        "currency_code": p.currency_code,
        "currency_symbol": p.currency_symbol,
        "min_price": p.min_price,
        "max_price": p.max_price,
        "price_type": p.price_type,
        "denominations": p.denominations,
        "thumbnail": p.thumbnail,
        "mobile_image": p.mobile_image,
        "base_image": p.base_image,
        "small_image": p.small_image,
        "tnc_link": p.tnc_link,
        "tnc_content": p.tnc_content,
        "expiry_info": p.expiry_info,
        "in_stock": p.in_stock,
        "is_digital": p.is_digital,
        "kyc_enabled": p.kyc_enabled,
        "disable_cart": p.disable_cart,
        "provider_created_at": p.provider_created_at,
        "provider_updated_at": p.provider_updated_at,
        "last_list_synced_at": p.last_list_synced_at,
        "last_detail_synced_at": p.last_detail_synced_at,
        "created_at": p.created_at,
        "updated_at": p.updated_at,
    }


def _override_row(o: ProductOverride) -> dict:
    logo = ""
    if o.brand_logo:
        try:
            logo = o.brand_logo.url
        except Exception:
            logo = o.brand_logo.name or ""
    return {
        "override_id": o.pk,
        "product_id": o.product_id,
        "sku": o.product.sku,
        "product_name": o.product.name,
        "category_id": o.category_id,
        "category_name": o.category.name if o.category_id else "",
        "name": o.name,
        "description": o.description,
        "tnc_content": o.tnc_content,
        "margin": o.margin,
        "status": o.status,
        "status_reason": o.status_reason,
        "brand_name": o.brand_name,
        "brand_code": o.brand_code,
        "brand_logo": logo,
        "logo_domain": o.logo_domain,
        "logo_dev_etag": o.logo_dev_etag,
        "logo_fetched_at": o.logo_fetched_at,
        "thumbnail": o.thumbnail,
        "created_at": o.created_at,
        "updated_at": o.updated_at,
    }


def _load_rates_csv(path: Path) -> dict[str, Decimal]:
    rates = {}
    if not path.is_file():
        return rates
    with path.open(newline="", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)
        for row in reader:
            sku = (row.get("SKU ID") or row.get("SKU") or "").strip()
            if not sku:
                continue
            raw = (
                row.get("UPI discount %")
                or row.get("UPI discount (% off face value)")
                or ""
            ).strip().replace("%", "")
            if not raw:
                continue
            rates[sku] = Decimal(raw)
    return rates


def _write_sheet(ws, rows: list[dict]):
    if not rows:
        ws.append(["(no rows)"])
        return
    headers = list(rows[0].keys())
    ws.append(headers)
    for row in rows:
        ws.append([_cell(row.get(h)) for h in headers])


class Command(BaseCommand):
    help = "Export gift card products, overrides, and margin sheet reconciliation to Excel."

    def add_arguments(self, parser):
        parser.add_argument(
            "--output",
            type=str,
            default="",
            help="Output .xlsx path (default: exports/giftcard_catalog_<timestamp>.xlsx)",
        )
        parser.add_argument(
            "--rates-csv",
            type=str,
            default=str(DEFAULT_RATES_CSV),
            help="Discount sheet CSV for margin reconciliation sheet",
        )

    def handle(self, *args, **options):
        try:
            from openpyxl import Workbook
        except ImportError as exc:
            self.stderr.write(
                "openpyxl is required. Install with: pip install openpyxl"
            )
            raise SystemExit(1) from exc

        out_path = options["output"].strip()
        if not out_path:
            stamp = timezone.now().strftime("%Y%m%d_%H%M%S")
            out_path = f"exports/giftcard_catalog_{stamp}.xlsx"
        out_path = Path(out_path)
        out_path.parent.mkdir(parents=True, exist_ok=True)

        products = list(
            ProviderProduct.objects.select_related("provider").order_by("sku")
        )
        overrides = list(
            ProductOverride.objects.select_related("product", "category").order_by(
                "product__sku"
            )
        )

        product_rows = [_product_row(p) for p in products]
        override_rows = [_override_row(o) for o in overrides]

        rates = _load_rates_csv(Path(options["rates_csv"]))
        db_by_sku = {p.sku: p for p in products}
        override_by_sku = {o.product.sku: o for o in overrides}

        matched_rows = []
        missing_db_rows = []
        for sku, sheet_margin in sorted(rates.items()):
            prod = db_by_sku.get(sku)
            if not prod:
                missing_db_rows.append(
                    {
                        "sku": sku,
                        "sheet_upi_discount_pct": sheet_margin,
                        "status": "In sheet, not in catalog",
                    }
                )
                continue
            ov = override_by_sku.get(sku)
            db_margin = ov.margin if ov and ov.margin is not None else ""
            matched_rows.append(
                {
                    "sku": sku,
                    "product_name": prod.name,
                    "in_stock": prod.in_stock,
                    "sheet_upi_discount_pct": sheet_margin,
                    "db_margin": db_margin,
                    "margin_matches_sheet": (
                        ov is not None
                        and ov.margin is not None
                        and Decimal(str(ov.margin)) == sheet_margin
                    ),
                }
            )

        in_stock_not_in_sheet = []
        sheet_skus = set(rates.keys())
        for p in products:
            if not p.in_stock:
                continue
            if p.sku in sheet_skus:
                continue
            ov = override_by_sku.get(p.sku)
            in_stock_not_in_sheet.append(
                {
                    "sku": p.sku,
                    "product_name": p.name,
                    "in_stock": True,
                    "db_margin": ov.margin if ov and ov.margin is not None else "",
                    "status": "In catalog (in stock), not in rate sheet",
                }
            )
        in_stock_not_in_sheet.sort(key=lambda r: r["sku"])

        updated_count = sum(1 for r in matched_rows if r["margin_matches_sheet"])

        wb = Workbook()
        ws_products = wb.active
        ws_products.title = "Products"
        _write_sheet(ws_products, product_rows)

        ws_over = wb.create_sheet("ProductOverrides")
        _write_sheet(ws_over, override_rows)

        ws_margin = wb.create_sheet("Margin sheet reconciliation")
        ws_margin.append(["Metric", "Value"])
        ws_margin.append(["Total products in DB", len(products)])
        ws_margin.append(["Total product overrides in DB", len(overrides)])
        ws_margin.append(["SKUs in rate sheet", len(rates)])
        ws_margin.append(["SKUs in sheet and in catalog", len(matched_rows)])
        ws_margin.append(["Overrides with margin matching sheet", updated_count])
        ws_margin.append(["SKUs in sheet but not in catalog", len(missing_db_rows)])
        ws_margin.append(
            ["In-stock catalog SKUs not in sheet", len(in_stock_not_in_sheet)]
        )
        ws_margin.append([])
        ws_margin.append(["--- Matched (sheet + catalog) ---"])
        _write_sheet(ws_margin, matched_rows)
        ws_margin.append([])
        ws_margin.append(["--- In sheet, not in catalog ---"])
        _write_sheet(ws_margin, missing_db_rows)
        ws_margin.append([])
        ws_margin.append(["--- In-stock, not in sheet ---"])
        _write_sheet(ws_margin, in_stock_not_in_sheet)

        wb.save(out_path)
        self.stdout.write(self.style.SUCCESS(f"Wrote {out_path.resolve()}"))
